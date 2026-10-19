"""
Демо-парсер: собирает все книги с учебного сайта books.toscrape.com
и сохраняет их в Excel-таблицу с оформлением.

Сайт books.toscrape.com создан специально для тренировки парсинга,
собирать с него данные можно свободно.
"""
import time
from datetime import datetime
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

START_URL = 'https://books.toscrape.com/catalogue/page-1.html'
OUTPUT_FILE = 'books.xlsx'
DELAY = 0.5           # пауза между запросами (сек), чтобы не нагружать сайт
MAX_PAGES = 100       # ограничение на всякий случай

# Рейтинг на сайте записан словом в классе: "star-rating Three"
RATINGS = {'One': 1, 'Two': 2, 'Three': 3, 'Four': 4, 'Five': 5}

HEADERS = {'User-Agent': 'Mozilla/5.0 (demo parser for learning)'}


def get_page(session, url):
    """Скачивает страницу. При ошибке сети пробует ещё 2 раза."""
    for attempt in range(3):
        try:
            response = session.get(url, headers=HEADERS, timeout=15)
            response.raise_for_status()
            response.encoding = 'utf-8'
            return response.text
        except requests.RequestException as e:
            print(f'  Ошибка: {e}. Повтор через 3 сек...')
            time.sleep(3)
    raise RuntimeError(f'Не удалось загрузить {url}')


def parse_books(html, page_url):
    """Достаёт из страницы список книг и ссылку на следующую страницу."""
    soup = BeautifulSoup(html, 'html.parser')
    books = []

    for card in soup.select('article.product_pod'):
        link = card.select_one('h3 a')
        title = link.get('title', link.get_text(strip=True))
        url = urljoin(page_url, link['href'])

        price_text = card.select_one('p.price_color').get_text(strip=True)
        # '£51.77' -> 51.77 (оставляем только цифры и точку)
        price = float(''.join(ch for ch in price_text if ch.isdigit() or ch == '.'))

        rating_tag = card.select_one('p.star-rating')
        rating_word = next((c for c in rating_tag['class'] if c in RATINGS), None)
        rating = RATINGS.get(rating_word, 0)

        stock = card.select_one('p.instock.availability')
        in_stock = 'Да' if stock and 'In stock' in stock.get_text() else 'Нет'

        books.append({'title': title, 'price': price, 'rating': rating,
                      'in_stock': in_stock, 'url': url})

    next_link = soup.select_one('li.next a')
    next_url = urljoin(page_url, next_link['href']) if next_link else None
    return books, next_url


def save_to_excel(books, filename):
    """Сохраняет книги в Excel: заголовки, фильтры, ширина колонок, ссылки."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Книги'

    headers = ['№', 'Название', 'Цена, £', 'Рейтинг', 'В наличии', 'Ссылка']
    ws.append(headers)

    for i, b in enumerate(books, start=1):
        ws.append([i, b['title'], b['price'], '★' * b['rating'], b['in_stock'], b['url']])
        cell = ws.cell(row=i + 1, column=6)
        cell.hyperlink = b['url']
        cell.font = Font(color='0563C1', underline='single')
        ws.cell(row=i + 1, column=3).number_format = '0.00'

    # Оформление шапки
    header_fill = PatternFill('solid', fgColor='1F4E78')
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    widths = [7, 60, 13, 14, 15, 50]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = 'A2'                 # шапка не уезжает при прокрутке
    ws.auto_filter.ref = ws.dimensions     # фильтры по всем колонкам

    # Второй лист — короткая сводка
    summary = wb.create_sheet('Сводка')
    prices = [b['price'] for b in books] or [0]
    rows = [
        ('Дата сбора', datetime.now().strftime('%d.%m.%Y %H:%M')),
        ('Всего книг', len(books)),
        ('Средняя цена, £', round(sum(prices) / len(prices), 2)),
        ('Самая дешёвая, £', min(prices)),
        ('Самая дорогая, £', max(prices)),
        ('С рейтингом 5 звёзд', sum(1 for b in books if b['rating'] == 5)),
    ]
    for row in rows:
        summary.append(row)
    for cell in summary['A']:
        cell.font = Font(bold=True)
    summary.column_dimensions['A'].width = 24
    summary.column_dimensions['B'].width = 20

    wb.save(filename)


def main():
    print('Начинаю сбор данных с books.toscrape.com...')
    session = requests.Session()
    all_books = []
    url = START_URL
    page = 0

    while url and page < MAX_PAGES:
        page += 1
        print(f'Страница {page}: {url}')
        html = get_page(session, url)
        books, url = parse_books(html, url)
        all_books.extend(books)
        print(f'  найдено книг: {len(books)}, всего: {len(all_books)}')
        time.sleep(DELAY)

    save_to_excel(all_books, OUTPUT_FILE)
    print(f'\nГотово! Собрано книг: {len(all_books)}. Файл: {OUTPUT_FILE}')


if __name__ == '__main__':
    main()
